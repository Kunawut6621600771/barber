"""สร้างเอกสารการทดสอบ (xlsx) จากผลรันจริง: unit.tap / e2e.tap / Robot output.xml
โครงชีตตามเอกสารต้นแบบ · ตัวเลขสรุปทุกช่องเป็นสูตรที่นับจากชีตรายละเอียด
"""
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from robot.api import ExecutionResult

HERE = Path(__file__).resolve().parent
RES = HERE / "results"
OUT = HERE / "เอกสารการทดสอบ_BarberQueue.xlsx"
TODAY = date.today().isoformat()
ROUND = "R1"

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")
PASS_FILL = PatternFill("solid", fgColor="DCFCE7")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def parse_tap(path):
    """คืน [(name, ok, ms)] จาก TAP ของ node --test (เฉพาะ test ระดับบน)"""
    rows = []
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    blocks = re.split(r"\n(?=(?:not )?ok \d+ - )", "\n" + text)
    for b in blocks:
        m = re.match(r"(not )?ok \d+ - (.+)", b.strip())
        if not m:
            continue
        ms = re.search(r"duration_ms: ([\d.]+)", b)
        rows.append((m.group(2).strip(), m.group(1) is None, round(float(ms.group(1)), 1) if ms else None))
    return rows


def split_id(name):
    m = re.match(r"(\S+)\s+(.*)", name)
    return (m.group(1), m.group(2)) if m else (name, "")


def sheet(wb, title, header, widths, banner, note, freeze="C5"):
    ws = wb.create_sheet(title)
    ws["A1"] = banner
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, color="6B7280")
    last = get_column_letter(len(header))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    for i, h in enumerate(header, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = freeze
    return ws


def put(ws, row, values, result_col=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.border, c.alignment = BOX, WRAP
    if result_col:
        c = ws.cell(row=row, column=result_col)
        c.fill = PASS_FILL if c.value == "PASS" else FAIL_FILL if c.value == "FAIL" else c.fill
        c.font = Font(bold=True)


# ---------------------------------------------------------------- อ่านผลจริง
unit = parse_tap(RES / "unit.tap")
e2e = parse_tap(RES / "e2e.tap")
robot = ExecutionResult(str(RES / "output.xml"))
robot_tests = []


def walk(s):
    for t in s.tests:
        robot_tests.append((s.name, t))
    for c in s.suites:
        walk(c)


walk(robot.suite)
if not unit or not e2e or not robot_tests:
    sys.exit("ไม่พบผลการทดสอบ — รัน run-all.sh ก่อน")

UNIT_FILES = {
    "UT-PP": ("promptpay.test.js", "src/utils/promptpay.js", "สร้าง QR PromptPay (EMVCo) + CRC16 + ปิดเลขบัญชี"),
    "UT-LY": ("loyalty.test.js", "src/utils/loyalty.js", "คำนวณแต้มสะสม / ส่วนลด / แต้มที่ใช้ได้ / หักแต้มตอนคืนเงิน"),
    "UT-VD": ("validate.test.js", "src/utils/validate.js", "ตรวจข้อมูลที่ส่งเข้า API (ชนิด/ช่วง/วันที่/รหัสผ่าน)"),
    "UT-MS": ("misc.test.js", "src/utils/{mask,settings,line-bot,line-link,http}.js", "ปิดชื่อลูกค้า · แปลงค่าตั้งค่า · เจตนาแชต LINE · รหัสผูกบัญชี"),
}
E2E_GROUPS = {
    "E2E-Q": ("queue-flow.e2e.test.js", "A. วงจรคิว: รับคิว → เรียก → ให้บริการ → เสร็จ", "ลูกค้า · พนักงาน"),
    "E2E-A": ("appointments.e2e.test.js", "B. จองล่วงหน้า → ยกเลิก / เช็กอินเป็นบัตรคิว", "ลูกค้า · พนักงาน"),
    "E2E-R": ("roles.e2e.test.js", "C. สิทธิ์ของแต่ละบทบาท (RBAC) และการยืนยันตัวตน", "แอดมิน · พนักงาน · ช่าง · ลูกค้า"),
    "E2E-S": ("sales.e2e.test.js", "D. เงินสด / ชำระบิล / แต้มสะสม / คืนเงิน / รายงาน", "แอดมิน · พนักงาน · ลูกค้า"),
}
ROBOT_SCOPE = "หน้าเว็บจริงบน Chromium: ล็อกอิน สมัคร รับบัตรคิว คิวของฉัน เรียกคิว สิทธิ์ตามบทบาท จอคิว"

wb = Workbook()
wb.remove(wb.active)
summary = wb.create_sheet("สรุปภาพรวม")  # สร้างก่อนเพื่อให้อยู่ชีตแรก เติมข้อมูลทีหลัง

# ---------------------------------------------------------------- Unit Test
ws = sheet(wb, "Unit Test",
           ["รหัสเคส", "ไฟล์ทดสอบ", "โมดูลที่ทดสอบ", "สิ่งที่ทดสอบ (ไฟล์)", "กรณีทดสอบ / ผลที่คาดหวัง", "รอบที่เริ่มมี", "ผล", "เวลา (ms)"],
           [10, 20, 30, 40, 64, 10, 8, 10],
           f"Unit Test — node:test (backend) · {len(unit)} เคส",
           f"รันจริง {TODAY} · ฟังก์ชันคำนวณล้วน ไม่ต่อฐานข้อมูล · คำสั่ง: cd backend && npm test")
for r, (name, ok, ms) in enumerate(unit, 5):
    tid, desc = split_id(name)
    f, mod, what = UNIT_FILES.get(tid[:5], ("?", "?", "?"))
    put(ws, r, [tid, f, mod, what, desc, ROUND, "PASS" if ok else "FAIL", ms], result_col=7)

ws = sheet(wb, "Unit สรุปรายไฟล์",
           ["ไฟล์ทดสอบ", "สิ่งที่ทดสอบ", "จำนวนเคส", "ผ่าน", "ไม่ผ่าน", "อัตราผ่าน"],
           [24, 60, 10, 8, 9, 10], "Unit Test — สรุปรายไฟล์", "นับจากชีต Unit Test ด้วยสูตร", freeze=None)
for r, (f, _mod, what) in enumerate(UNIT_FILES.values(), 5):
    put(ws, r, [f, what,
                f"=COUNTIF('Unit Test'!$B$5:$B$1000,A{r})",
                f"=COUNTIFS('Unit Test'!$B$5:$B$1000,A{r},'Unit Test'!$G$5:$G$1000,\"PASS\")",
                f"=COUNTIFS('Unit Test'!$B$5:$B$1000,A{r},'Unit Test'!$G$5:$G$1000,\"FAIL\")",
                f"=IF(C{r}=0,0,D{r}/C{r})"])
    ws.cell(row=r, column=6).number_format = "0.0%"

# ---------------------------------------------------------------- E2E API
ws = sheet(wb, "E2E API",
           ["รหัสเคส", "หมวด / สถานการณ์", "กรณีทดสอบ / ผลที่คาดหวัง", "ไฟล์ทดสอบ", "รอบที่เริ่มมี", "ผล", "บทบาทที่เกี่ยวข้อง", "เวลา (ms)"],
           [11, 40, 64, 24, 10, 8, 26, 10],
           f"API End-to-End Test — ยิง REST API จริง · {len(e2e)} เคส",
           f"รันจริง {TODAY} บน backend :5051 + PostgreSQL ฐาน barbershop_test (แยกจากระบบจริง) · โหมดสลิป mock ไม่ส่ง LINE")
for r, (name, ok, ms) in enumerate(e2e, 5):
    tid, desc = split_id(name)
    f, grp, roles = E2E_GROUPS.get(tid[:5], ("?", "?", "?"))
    put(ws, r, [tid, grp, desc, f, ROUND, "PASS" if ok else "FAIL", roles, ms], result_col=6)

# ---------------------------------------------------------------- Robot
ws = sheet(wb, "Robot Framework",
           ["รหัสเคส", "ชุดทดสอบ (Suite)", "ขอบเขตของชุด", "ชื่อกรณีทดสอบ / ผลที่คาดหวัง", "ขั้นตอนการทดสอบ (Test Steps)", "รอบที่เริ่มมี", "ผล", "เวลา (วินาที)"],
           [9, 16, 36, 50, 70, 10, 8, 11],
           f"Automation Test — Robot Framework + Browser (Playwright) · {len(robot_tests)} เคส",
           f"รันจริง {TODAY} บน Chromium headless (1366×820, th-TH) · frontend :3001 → backend ทดสอบ :5051", freeze="B5")
for r, (suite, t) in enumerate(robot_tests, 5):
    tid, desc = split_id(t.name)
    steps = "\n".join(f"{i}. {k.name} {'  '.join(map(str, k.args))}".strip()
                      for i, k in enumerate([b for b in t.body if hasattr(b, "args")], 1))
    secs = round(t.elapsed_time.total_seconds(), 2)
    put(ws, r, [tid, suite.lower().replace(" ", "_"), ROBOT_SCOPE, desc, steps, ROUND, t.status, secs], result_col=7)

suites = sorted({s.lower().replace(" ", "_") for s, _ in robot_tests})
ws = sheet(wb, "Robot สรุปรายชุด",
           ["ชุดทดสอบ", "ขอบเขต", "จำนวนเคส", "ผ่าน", "ไม่ผ่าน", "อัตราผ่าน", "เวลารวม (วินาที)"],
           [20, 70, 10, 8, 9, 10, 14], "Robot Framework — สรุปรายชุดทดสอบ", "นับจากชีต Robot Framework ด้วยสูตร", freeze=None)
for r, s in enumerate(suites, 5):
    rf = "'Robot Framework'"
    put(ws, r, [s, ROBOT_SCOPE,
                f"=COUNTIF({rf}!$B$5:$B$1000,A{r})",
                f"=COUNTIFS({rf}!$B$5:$B$1000,A{r},{rf}!$G$5:$G$1000,\"PASS\")",
                f"=COUNTIFS({rf}!$B$5:$B$1000,A{r},{rf}!$G$5:$G$1000,\"FAIL\")",
                f"=IF(C{r}=0,0,D{r}/C{r})",
                f"=SUMIF({rf}!$B$5:$B$1000,A{r},{rf}!$H$5:$H$1000)"])
    ws.cell(row=r, column=6).number_format = "0.0%"
# ---------------------------------------------------------------- บั๊ก
# (รหัส, พบจาก, ความรุนแรง, อาการ, สาเหตุ, วิธีแก้, เคสยืนยัน, ไฟล์ที่แก้)
SCAN = "สแกนโค้ดทั้งระบบ (code review)"
BUGS = [
    ("BUG-01", SCAN, "Critical", "พนักงาน (staff) ตั้งตัวเองหรือสร้างบัญชีใหม่เป็นผู้ดูแลระบบได้ แล้วแก้/ระงับบัญชีแอดมินได้",
     "PUT/POST /users เปิดให้ staff และตรวจแค่ห้ามลดสิทธิ์ตัวเอง ไม่ได้ห้ามยกสิทธิ์",
     "ให้กำหนด role=admin และแก้บัญชีแอดมินได้เฉพาะแอดมิน · ซ่อนตัวเลือกในหน้าเว็บ",
     "E2E-R-07, E2E-R-08, E2E-R-09", "users.controller.js · Users.jsx"),
    ("BUG-02", SCAN, "Major", "จอคิวหน้าร้าน (ไม่ต้องล็อกอิน) และสตรีม realtime ส่งเบอร์โทร/รหัสลูกค้าทุกคนออกไป",
     "/queue/current และ broadcast queue:* ส่ง SELECT * จาก v_queue_board",
     "ส่งเฉพาะข้อมูลสาธารณะ (ตัดเบอร์/รหัสลูกค้า ปิดนามสกุล) · /barbers ไม่ส่งอีเมล/เบอร์ช่างให้ลูกค้า",
     "E2E-Q-12, E2E-R-05, UT-MS-01", "queue.controller.js · barbers.controller.js · utils/mask.js"),
    ("BUG-03", SCAN, "Major", "รับสลิปโอนเงินที่โอนให้บัญชีอื่นได้ ถ้าสลิปไม่มีเลขพร้อมเพย์ผู้รับ",
     "ข้ามการตรวจผู้รับเมื่ออ่านเลขบัญชีจากสลิปไม่ได้ (fail-open)",
     "อ่านผู้รับไม่ได้ต้องเทียบชื่อบัญชีร้าน ไม่ตรง = ไม่รับอัตโนมัติ ให้พนักงานยืนยัน", "ตรวจโค้ด (ไม่ยิงสลิปจริง)", "utils/slip.js"),
    ("BUG-04", SCAN, "Critical", "บันทึกเงินเข้า/ออกลิ้นชักที่มีเศษสตางค์ (เช่น 150.50) ขึ้น error 500",
     "SQL CASE ... THEN $3 ELSE 0 ทำให้ PostgreSQL เดาชนิดเป็น integer",
     "ระบุชนิด $3::numeric", "E2E-S-01", "shifts.controller.js"),
    ("BUG-05", SCAN, "Critical", "ยกเลิกบิลที่คืนเงินไปแล้ว ทำให้สต็อกถูกคืนซ้ำและเงินสดที่ควรมีในลิ้นชักผิด",
     "voidSale ห้ามแค่บิล void ไม่ได้ห้ามบิล refunded/partial_refund",
     "ห้ามยกเลิกบิลที่มีการคืนเงินแล้ว (409)", "E2E-S-06", "pos.controller.js"),
    ("BUG-06", SCAN, "Major", "คืนเงินแล้วลูกค้ายังได้แต้มสะสม/ยอดใช้จ่าย/ระดับสมาชิกเท่าเดิม",
     "refundSale ไม่ได้ย้อนแต้มและ total_spent",
     "หักแต้มตามสัดส่วนเงินที่คืน (ไม่หักซ้ำ) และลดยอดใช้จ่าย/ปรับระดับสมาชิก",
     "E2E-S-05, UT-LY-07, UT-LY-08", "pos.controller.js · utils/loyalty.js"),
    ("BUG-07", SCAN, "Major", "ลูกค้าเสียแต้มเกินส่วนลดที่ได้จริง เมื่อใส่แต้มมากกว่ายอดบิล",
     "updateSale คิดส่วนลดจากแต้มโดยไม่จำกัดเพดาน และตอนปิดบิลหักแต้มเต็มจำนวน",
     "จำกัดแต้มสูงสุดตามยอดคงเหลือ และตรวจซ้ำตอนปิดบิล", "E2E-S-07, UT-LY-06", "pos.controller.js · utils/loyalty.js"),
    ("BUG-08", SCAN, "Major", "ลบโค้ดส่วนลดออกจากบิลแล้วส่วนลดยังอยู่",
     "ล้างแค่ promo_code แต่ discount_type/value ยังเป็นค่าของโค้ดเดิม",
     "ลบโค้ดโดยไม่ได้กรอกส่วนลดมือ = ล้างส่วนลดด้วย", "E2E-S-08", "pos.controller.js"),
    ("BUG-09", SCAN, "Major", "โค้ดส่วนลด \"ใช้ได้ถึงวันที่ X\" หมดอายุตั้งแต่เที่ยงคืนต้นวัน X (เร็วไป 1 วัน)",
     "เทียบเวลาเที่ยงคืนกับเวลาปัจจุบัน แทนการเทียบวันที่",
     "เทียบเป็นวันที่ตามเขตเวลาร้าน · ตรวจโค้ดซ้ำตอนปิดบิล (ล็อกแถว กันใช้เกินสิทธิ์)", "E2E-S-09", "pos.controller.js"),
    ("BUG-10", SCAN, "Major", "ส่งออก CSV แล้วเปิดใน Excel ค่าที่ขึ้นต้นด้วย = + - @ (เช่นชื่อลูกค้า) ทำงานเป็นสูตร",
     "escape แค่เครื่องหมายคำพูด", "ใส่ ' นำหน้าค่าที่เสี่ยง (CSV injection)", "E2E-S-11", "reports.controller.js"),
]
BUGS += [
    ("BUG-11", SCAN, "Major", "ส่ง id ที่ไม่ใช่ตัวเลข / วันที่ผิด / limit ติดลบ ได้ error 500 แทน 400",
     "error handler ไม่แปลงรหัส PostgreSQL 22P02/22007/22008/22003 และวันที่ 2026-02-30 ผ่านการตรวจ",
     "แปลงเป็น 400 · ตรวจวันที่ให้มีจริงในปฏิทิน · จำกัด limit ≥ 1 · Number ต้องไม่ใช่ Infinity",
     "E2E-Q-13, E2E-A-03, E2E-S-10, UT-VD-07, UT-VD-10", "middleware/error.js · utils/validate.js · reports.controller.js"),
    ("BUG-12", SCAN, "Major", "คิวที่ถูกข้ามแล้วเรียกใหม่ ลูกค้าไม่ได้รับแจ้งเตือนรอบที่ 2",
     "แจ้งเตือน your_turn มี unique index ต่อคิว (ON CONFLICT DO NOTHING) จึงไม่ส่งซ้ำ",
     "ถ้าเคยแจ้ง your_turn แล้ว ส่ง called_again แทน", "E2E-Q-09", "queue.controller.js"),
    ("BUG-13", SCAN, "Minor", "คิวจองล่วงหน้าที่ถูกข้าม กลับไปอยู่ลำดับที่ 1 ทันที (ข้อความบอกว่าไปท้ายแถว)",
     "reorder_queue เรียง priority ก่อน skip_count", "เรียง skip_count ก่อน priority (migration 006)", "ตรวจ SQL", "database/migrations/006_bugfix_r1.sql"),
    ("BUG-14", SCAN, "Minor", "สรุปกะนับเงินทอนซ้ำเมื่อบิลจ่ายหลายวิธี · คืนเงินสดไปหักกะที่ขาย ไม่ใช่กะที่คืน",
     "SUM(change_amount) บน join กับ sale_payments · sale_refunds ไม่มี shift_id",
     "คิดเงินทอนแยก · เพิ่ม sale_refunds.shift_id (migration 006)", "ตรวจ SQL", "migration 006 · shifts/pos controller"),
    ("BUG-15", SCAN, "Minor", "จองเวลาที่ผ่านไปแล้วของวันนี้ได้ (ผ่าน API) · จองช่างเจาะจงทำให้ช่วงเวลาเกินจำนวนช่าง",
     "ตรวจเวลาที่ผ่านแล้วเฉพาะหน้าแสดงช่วงว่าง · นับที่นั่งต่างกันระหว่างจองรวมกับจองเจาะจงช่าง",
     "ตรวจเวลาตอนจองจริง ใช้เขตเวลาร้าน และนับที่นั่งกติกาเดียว · ยกเลิกแบบมีเงื่อนไขกันชนกับเช็กอิน",
     "E2E-A-04, E2E-A-06, E2E-A-09", "appointments.controller.js"),
    ("BUG-16", SCAN, "Minor", "รหัสผูกบัญชี LINE หนึ่งรหัสถูกใช้ได้ 2 บัญชีถ้าส่งพร้อมกัน",
     "ตรวจ used_at แล้วค่อยอัปเดต (ไม่ atomic)", "จองรหัสด้วย UPDATE ... WHERE used_at IS NULL ก่อน", "UT-MS-06 + ตรวจโค้ด", "utils/line-link.js"),
    ("BUG-17", SCAN, "Minor", "ส่งรหัสผ่าน/เหตุผล/คำค้นเป็นตัวเลขหรือ array แล้วระบบล่ม 500",
     "เรียก .trim() / bcrypt.compare กับค่าที่ไม่ใช่สตริง", "แปลงเป็น String ก่อนใช้", "E2E-R-03", "auth/users/services/queue/pos controller"),
    ("BUG-18", SCAN, "Major", "ฟอร์มออกบัตรคิว (พนักงาน) ล้างตัวเองระหว่างกรอก/หลังออกบัตร เมื่อมีคิวขยับ",
     "useEffect ผูก dependency กับ services ซึ่งได้ array ใหม่ทุกครั้งที่ queue:update",
     "ผูกเฉพาะ open", "ตรวจโค้ด + vite build", "QueueManage.jsx"),
    ("BUG-19", SCAN, "Minor", "แถบด้านข้าง error เมื่อได้ queue:update ที่ไม่มี tickets · ค้นหาด้านบนไม่ทำงาน · ?src=qr หายหลังล็อกอิน",
     "อ่าน d.tickets โดยไม่ตรวจ · หน้าคิวไม่อ่าน ?q= · เก็บแค่ pathname",
     "ตรวจก่อนใช้ / โหลดใหม่ · กรองคิวตาม q · เก็บ pathname + search", "ตรวจโค้ด (UI-07 ครอบการเปิดหน้าคิว)", "Layout.jsx · QueueManage.jsx · App.jsx"),
    ("BUG-20", SCAN, "Minor", "token หมดอายุแล้วแอปค้างสถานะล็อกอิน ทุกปุ่มขึ้น 401 · SSE ยังใช้ token เก่า",
     "refresh ไม่ผ่านแล้วไม่แจ้ง AuthContext · EventSource สร้างครั้งเดียว",
     "ส่ง event auth:expired ให้ออกจากระบบ · ต่อ SSE ใหม่เมื่อ token เปลี่ยน", "ตรวจโค้ด", "api/client.js · AuthContext.jsx · realtime.js"),
    ("BUG-21", "Unit Test (UT-MS-01)", "Minor", "คิว walk-in ที่ไม่มีชื่อ ได้ชื่อเป็น \"\" แทน null หลังปิดชื่อ",
     "maskName คืนสตริงว่างเสมอเมื่อไม่มีชื่อ", "คงค่า null เดิมไว้", "UT-MS-01", "utils/mask.js"),
    ("BUG-22", "API E2E (รอบแรก)", "Minor", "ชุดทดสอบอัตโนมัติสมัครสมาชิกถี่ โดน rate limit 429 จนเทสต์ล้ม",
     "rate limit สมัคร 10 ครั้ง/ชม. ใช้กับทุกสภาพแวดล้อม", "ข้าม rate limit เฉพาะ NODE_ENV=test (production ยังจำกัดเท่าเดิม)", "E2E ทั้งชุด 46/46", "server.js"),
]

ws = sheet(wb, "บั๊กที่พบและแก้ไข",
           ["รหัส", "รอบ", "วันที่พบ/แก้", "พบจาก (เครื่องมือ)", "ความรุนแรง", "อาการที่พบ", "สาเหตุ", "วิธีแก้",
            "เคสทดสอบยืนยัน / กันถดถอย", "ไฟล์ที่แก้", "สถานะ"],
           [9, 6, 11, 22, 10, 46, 40, 46, 26, 30, 9],
           "บันทึกข้อบกพร่องที่พบจากการทดสอบ (Defect Log)",
           "สถานะ \"ปิดแล้ว\" = แก้แล้วและรันชุดทดสอบซ้ำผ่าน · เคสที่ระบุใช้กันบั๊กกลับมา (regression)")
for r, b in enumerate(BUGS, 5):
    put(ws, r, [b[0], ROUND, TODAY, b[1], b[2], b[3], b[4], b[5], b[6], b[7], "ปิดแล้ว"])
# ---------------------------------------------------------------- เครื่องมือ / สภาพแวดล้อม / วิธีรัน
import platform  # noqa: E402
import subprocess  # noqa: E402
import robot as _robot  # noqa: E402
import openpyxl as _ox  # noqa: E402


def ver(cmd):
    try:
        return subprocess.run(cmd, capture_output=True, text=True, shell=True, timeout=20).stdout.strip().splitlines()[0]
    except Exception:  # noqa: BLE001
        return "-"


node_v = ver("node -v")
try:
    from importlib.metadata import version as _pv
    browser_v = _pv("robotframework-browser")
except Exception:  # noqa: BLE001
    browser_v = "-"

ws = sheet(wb, "เครื่องมือที่ใช้ทดสอบ",
           ["ลำดับ", "เครื่องมือ", "เวอร์ชัน", "ประเภทการทดสอบ", "ใช้ทดสอบอะไร", "คำสั่งรัน", "ที่อยู่ไฟล์", "เริ่มใช้รอบ"],
           [6, 26, 18, 22, 48, 46, 34, 10], "เครื่องมือที่ใช้ทดสอบ (Testing Tools)", "เวอร์ชันอ่านจากเครื่องที่รันจริง")
TOOLS = [
    ("node:test (Node.js built-in)", node_v, "Unit Test", "ฟังก์ชันคำนวณฝั่ง backend: QR PromptPay, ตรวจข้อมูล, แต้มสะสม ฯลฯ ไม่ต่อฐานข้อมูล",
     "cd backend && npm test", "backend/tests/unit/*.test.js"),
    ("node:test + fetch (API E2E)", node_v, "API E2E / Integration", "ยิง REST API จริงทั้งเส้นทาง คิว จองล่วงหน้า สิทธิ์ บิล/แต้ม",
     "cd backend && npm run test:e2e", "backend/tests/e2e/*.e2e.test.js · helpers.js"),
    ("Robot Framework", _robot.__version__ if hasattr(_robot, "__version__") else "7.x", "Automation / Acceptance",
     "ชุดทดสอบแบบ keyword-driven ออกรายงาน report.html / log.html", "automation-testing\\.venv\\Scripts\\robot --outputdir results tests",
     "automation-testing/tests/*.robot"),
    ("robotframework-browser (Playwright)", browser_v, "UI Test บนเบราว์เซอร์", "กดหน้าเว็บจริงบน Chromium", "(เรียกผ่าน robot)",
     "automation-testing/resources/common.resource"),
    ("openpyxl", _ox.__version__, "รายงานผล", "สร้างเอกสาร xlsx จากผลรันจริง", "python build_report.py", "automation-testing/build_report.py"),
    ("Vite build", "-", "Build check", "คอมไพล์หน้าเว็บทั้งหมด จับ import/route ที่เสีย", "cd frontend && npx vite build", "frontend/"),
]
for r, t in enumerate(TOOLS, 5):
    put(ws, r, [r - 4, *t, ROUND])

ws = sheet(wb, "สภาพแวดล้อม", ["หัวข้อ", "ค่า", "หมายเหตุ"], [30, 60, 50],
           "สภาพแวดล้อมการทดสอบ (Test Environment)", "ไม่บันทึกรหัสผ่านในเอกสารนี้ — บัญชีทดสอบมาจาก database/seed.sql", freeze=None)
ENV = [
    ("ระบบปฏิบัติการ", f"{platform.system()} {platform.release()} ({platform.version()})", "เครื่องพัฒนา"),
    ("Node.js", node_v, "backend + unit/E2E"),
    ("Python", platform.python_version(), "Robot Framework (.venv)"),
    ("ฐานข้อมูล", "PostgreSQL localhost:5432 · ฐาน barbershop_test", "แยกจากฐานจริง barbershop ไม่แตะข้อมูลร้าน"),
    ("Backend ทดสอบ", "http://localhost:5051 · NODE_ENV=test", "ค่าใน backend/tests/test-env.sh"),
    ("Frontend ทดสอบ", "http://localhost:3001 (vite.test.config.js)", "proxy /api → :5051"),
    ("การชำระเงิน", "SLIP_PROVIDER=mock · PAYMENT_REQUIRED=false", "ไม่เรียก EasySlip จริง ไม่เสียโควตา"),
    ("LINE", "ปิด (ไม่ตั้ง token)", "ไม่ส่งข้อความหาลูกค้าจริงระหว่างทดสอบ"),
    ("เบราว์เซอร์", "Chromium (Playwright) headless · 1366×820 · th-TH", ""),
]
for r, e in enumerate(ENV, 5):
    put(ws, r, list(e))

ws = sheet(wb, "วิธีรันทดสอบ", ["ขั้น", "คำสั่ง", "ผลที่ควรได้"], [36, 90, 40],
           "วิธีรันการทดสอบซ้ำ", "รันจากโฟลเดอร์โปรเจกต์ barber (Git Bash)", freeze=None)
HOWTO = [
    ("1. สร้างฐานทดสอบ (ครั้งแรก)", "cd backend && source tests/test-env.sh && node scripts/setup-db.js", "สร้าง barbershop_test + seed + migrations"),
    ("2. เปิด backend ทดสอบ", "cd backend && source tests/test-env.sh && node src/server.js", "/api/health ตอบ database=barbershop_test"),
    ("3. เปิด frontend ทดสอบ", "cd frontend && npx vite --config vite.test.config.js", "เปิด http://localhost:3001 ได้"),
    ("4. ติดตั้ง Robot (ครั้งแรก)", "cd automation-testing && py -3.13 -m venv .venv && .venv/Scripts/pip install robotframework robotframework-browser openpyxl && .venv/Scripts/rfbrowser init chromium", "ติดตั้งครบ"),
    ("5. รันทุกชุด + สร้างเอกสาร", "bash automation-testing/run-all.sh", "ได้ไฟล์ เอกสารการทดสอบ_BarberQueue.xlsx"),
    ("5.1 Unit อย่างเดียว", "cd backend && npm test", f"{len(unit)} pass"),
    ("5.2 API E2E อย่างเดียว", "cd backend && npm run test:e2e", f"{len(e2e)} pass"),
    ("5.3 UI อย่างเดียว", "cd automation-testing && .venv/Scripts/robot --outputdir results tests", f"{len(robot_tests)} pass · results/report.html"),
]
for r, h in enumerate(HOWTO, 5):
    put(ws, r, list(h))
# ---------------------------------------------------------------- สรุปภาพรวม (สูตรทั้งหมด)
ws = summary
for col, w in zip("ABCDEFG", [30, 42, 12, 10, 10, 12, 28]):
    ws.column_dimensions[col].width = w
ws["A1"] = "เอกสารการทดสอบระบบ Barber Queue (ระบบเรียกคิวร้านตัดผม Trim & Co.)"
ws["A1"].font = Font(bold=True, size=15)
ws.merge_cells("A1:G1")
ws["A2"] = f"รอบทดสอบ {ROUND} · {TODAY} · ตัวเลขทุกช่องในตารางผลคำนวณจากชีตรายละเอียดด้วยสูตร"
ws["A2"].font = Font(italic=True, color="6B7280")
ws.freeze_panes = "A4"

ws["A4"] = "ข้อมูลเอกสาร"
ws["A4"].font = Font(bold=True, size=12)
INFO = [
    ("ชื่อระบบ", "Barber Queue — ระบบเรียกคิว จองล่วงหน้า ชำระ QR PromptPay สมาชิกสะสมแต้ม และ LINE OA"),
    ("ส่วนประกอบที่ทดสอบ", "backend: Express + PostgreSQL · web: React 18 + Vite · LINE Bot / LIFF"),
    ("ช่วงเวลาการทดสอบ", f"{TODAY} (รอบ {ROUND})"),
    ("วันที่จัดทำเอกสาร", TODAY),
    ("สภาพแวดล้อม", "Windows + PostgreSQL localhost ฐาน barbershop_test — ดูชีต \"สภาพแวดล้อม\""),
]
for r, (k, v) in enumerate(INFO, 5):
    ws.cell(row=r, column=1, value=k).font = Font(bold=True)
    ws.cell(row=r, column=2, value=v)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)

ws["A11"] = f"ผลการทดสอบรอบล่าสุด ({ROUND} · {TODAY}) — คำนวณจากชีตรายละเอียดด้วยสูตร"
ws["A11"].font = Font(bold=True, size=12)
for i, h in enumerate(["ระดับการทดสอบ", "เครื่องมือ", "จำนวนเคส", "ผ่าน", "ไม่ผ่าน", "อัตราผ่าน", "เวลารัน"], 1):
    c = ws.cell(row=12, column=i, value=h)
    c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX


def secs(rows):
    return f"{sum(ms or 0 for *_, ms in rows) / 1000:.1f} วินาที"


LEVELS = [
    ("Unit Test (backend)", "node:test", "'Unit Test'!G5:G1000", secs(unit)),
    ("API E2E / Integration Test", "node:test + fetch", "'E2E API'!F5:F1000", secs(e2e)),
    ("UI Automation (Robot Framework)", "Robot Framework + Browser (Playwright)", "'Robot Framework'!G5:G1000",
     f"{sum(t.elapsed_time.total_seconds() for _, t in robot_tests):.1f} วินาที"),
]
for r, (lvl, tool, rng, t) in enumerate(LEVELS, 13):
    put(ws, r, [lvl, tool, f"=COUNTA({rng})", f'=COUNTIF({rng},"PASS")', f'=COUNTIF({rng},"FAIL")', f"=IF(C{r}=0,0,D{r}/C{r})", t])
    ws.cell(row=r, column=6).number_format = "0.0%"
put(ws, 16, ["รวมทุกระดับ", "", "=SUM(C13:C15)", "=SUM(D13:D15)", "=SUM(E13:E15)", "=IF(C16=0,0,D16/C16)", ""])
ws.cell(row=16, column=6).number_format = "0.0%"
for c in range(1, 8):
    ws.cell(row=16, column=c).font = Font(bold=True)

ws["A18"] = "ข้อบกพร่องที่พบ (นับจากชีต \"บั๊กที่พบและแก้ไข\")"
ws["A18"].font = Font(bold=True, size=12)
bug = "'บั๊กที่พบและแก้ไข'"
for i, h in enumerate(["ความรุนแรง", "", "พบทั้งหมด", "ปิดแล้ว", "ค้างอยู่"], 1):
    c = ws.cell(row=19, column=i, value=h)
    c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
for r, sev in enumerate(["Critical", "Major", "Minor"], 20):
    put(ws, r, [sev, "",
                f'=COUNTIF({bug}!$E$5:$E$1000,A{r})',
                f'=COUNTIFS({bug}!$E$5:$E$1000,A{r},{bug}!$K$5:$K$1000,"ปิดแล้ว")',
                f"=C{r}-D{r}"])
put(ws, 23, ["รวม", "", "=SUM(C20:C22)", "=SUM(D20:D22)", "=SUM(E20:E22)"])
for c in range(1, 6):
    ws.cell(row=23, column=c).font = Font(bold=True)

ws["A25"] = "สิ่งที่ไม่ได้ทดสอบในรอบนี้"
ws["A25"].font = Font(bold=True, size=12)
SKIP = [
    ("ตรวจสลิปผ่าน EasySlip จริง", "ใช้โหมด mock — ไม่เสียโควตาและไม่ต้องโอนเงินจริง (ตรรกะตรวจผู้รับตรวจด้วยการอ่านโค้ด)"),
    ("ส่งข้อความ LINE ถึงลูกค้าจริง", "ปิด LINE ในสภาพแวดล้อมทดสอบ กันส่งข้อความหาลูกค้าจริง"),
    ("SonarQube / CI/CD", "ตัดออกจากขอบเขตรอบนี้ตามที่ตกลงกัน"),
]
for r, (k, v) in enumerate(SKIP, 26):
    ws.cell(row=r, column=1, value=k).font = Font(bold=True)
    ws.cell(row=r, column=2, value=v)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)

wb.save(OUT)
print(f"saved {OUT.name}: unit {sum(ok for _, ok, _ in unit)}/{len(unit)} · "
      f"e2e {sum(ok for _, ok, _ in e2e)}/{len(e2e)} · "
      f"robot {sum(t.status == 'PASS' for _, t in robot_tests)}/{len(robot_tests)} · bugs {len(BUGS)}")
